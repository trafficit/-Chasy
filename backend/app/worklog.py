"""Time math and Excel import/export, ported from the desktop worklog_dashboard2.py."""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

HOLIDAY_MARK = "(Holiday)"


def parse_hhmm_to_minutes(value) -> int:
    try:
        h, m = str(value).split(":")
        return int(h) * 60 + int(m)
    except Exception:
        return 0


def minutes_to_hhmm(total: int) -> str:
    return f"{total // 60}:{total % 60:02d}"


def compute(date: str, start: str, end: str, lunch: str, comment: str) -> dict:
    """Return a fully computed entry dict, or raise ValueError with a user message."""
    date = (date or "").strip()
    comment = (comment or "").strip()

    if HOLIDAY_MARK in date:
        return {
            "date": date,
            "start": "00:00",
            "end": "00:00",
            "duration": "0:00",
            "lunch": "0:00",
            "net": "0:00",
            "comment": comment,
        }

    start = (start or "").strip()
    end = (end or "").strip()
    if ":" not in start:
        start += ":00"
    if ":" not in end:
        end += ":00"

    try:
        t1 = datetime.strptime(start, "%H:%M")
        t2 = datetime.strptime(end, "%H:%M")
    except ValueError:
        raise ValueError("Invalid time format (HH:MM)")

    total_minutes = int((t2 - t1).total_seconds() // 60)
    if total_minutes < 0:  # shift crosses midnight
        total_minutes += 24 * 60
    duration = minutes_to_hhmm(total_minutes)

    lunch = (lunch or "").strip()
    lunch_minutes = 0
    lunch_formatted = ""
    if lunch:
        try:
            if ":" in lunch:
                lh, lm = lunch.split(":")
                lunch_minutes = int(lh) * 60 + int(lm)
            else:
                lunch_minutes = int(lunch)
        except ValueError:
            raise ValueError("Lunch must be minutes or HH:MM")
        lunch_formatted = minutes_to_hhmm(lunch_minutes)

    net_minutes = max(total_minutes - lunch_minutes, 0)
    return {
        "date": date,
        "start": start,
        "end": end,
        "duration": duration,
        "lunch": lunch_formatted,
        "net": minutes_to_hhmm(net_minutes),
        "comment": comment,
    }


def build_xlsx(entries) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Worklog"

    lunch_present = any(e.lunch and "Holiday" not in e.date for e in entries)
    if lunch_present:
        headers = ["Date", "Start", "End", "Duration", "Lunch", "Net Duration", "Comment"]
    else:
        headers = ["Date", "Start", "End", "Duration", "Net Duration", "Comment"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for e in entries:
        if lunch_present:
            ws.append([e.date, e.start, e.end, e.duration, e.lunch, e.net, e.comment])
        else:
            ws.append([e.date, e.start, e.end, e.duration, e.net, e.comment])

    ws.append([])
    total_minutes = sum(parse_hhmm_to_minutes(e.net) for e in entries)
    total_str = f"{total_minutes // 60}:{total_minutes % 60:02d}"
    if lunch_present:
        ws.append(["▶ Net Total", "", "", "", "", total_str, ""])
    else:
        ws.append(["▶ Net Total", "", "", "", total_str, ""])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_xlsx(data: bytes) -> list[dict]:
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]

    def col(name: str):
        return header.index(name) if name in header else None

    i_date = col("date")
    i_start = col("start")
    i_end = col("end")
    i_lunch = col("lunch")
    i_comment = col("comment")

    def get(row, i):
        if i is None or i >= len(row) or row[i] is None:
            return None
        return row[i]

    out: list[dict] = []
    for row in rows[1:]:
        if not row:
            continue
        raw_date = get(row, i_date) if i_date is not None else (row[0] if row else None)
        if raw_date is None or str(raw_date).strip() == "":
            continue
        if str(raw_date).startswith("▶"):
            continue
        out.append(
            {
                "date": str(raw_date).strip(),
                "start": str(get(row, i_start) or "00:00"),
                "end": str(get(row, i_end) or "00:00"),
                "lunch": str(get(row, i_lunch) or ""),
                "comment": str(get(row, i_comment) or ""),
            }
        )
    return out
